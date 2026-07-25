from datetime import date

import pytest
from openpyxl import Workbook

from bot.storage import init_db, connect
from bot import investments_import as imp
from bot import investments_store as store

# NOTE: sqlite3 runs with detect_types=PARSE_DECLTYPES and storage.py:18-21
# registers DATE converters, so DATE columns come back as datetime.date,
# NOT str. Assertions against raw rows must use date(...) objects; only
# build_snapshot's payload is normalized to ISO strings via _as_iso().


def _make_xlsx(path, *, header_dates=("02-15-26", "07-25-26")):
    wb = Workbook()
    ws = wb.active
    ws.append(["Account", "Type", "Number", "Owner",
               f"2026 Value ({header_dates[0]})",
               f"2026 Value ({header_dates[1]})", "Notes"])
    ws.append(["Marcus", "Savings", "1234", "Joint", "$32,571.29", "$25,668.00", ""])
    ws.append(["117 Mayfield Dr", "Home Equity", "", "Joint",
               "$246,501.00", "$249,522.92", "Zestimate"])
    ws.append([])
    ws.append(["Type of Insurance", "Through Employer", "Provider", "Contact",
               "Coverage", "Deductible", "Annual Premium", "Comments", "Renewal"])
    ws.append(["Homeowners", "No", "Amica", "agent", "482k", "$1,000",
               "$2,798.00", "escrow", "2027-01-01"])
    wb.save(path)
    return path


def test_import_creates_rounds_holdings_values(tmp_path):
    db = tmp_path / "t.db"
    init_db(db)
    xlsx = _make_xlsx(tmp_path / "snap.xlsx")
    result = imp.import_xlsx(db, xlsx)
    assert result["rounds"] == 2
    assert result["holdings"] == 2
    assert result["values"] == 4
    rounds = store.list_rounds(db)
    assert [r["as_of_date"] for r in rounds] == [date(2026, 7, 25), date(2026, 2, 15)]


def test_import_is_idempotent(tmp_path):
    db = tmp_path / "t.db"
    init_db(db)
    xlsx = _make_xlsx(tmp_path / "snap.xlsx")
    imp.import_xlsx(db, xlsx)
    second = imp.import_xlsx(db, xlsx)
    assert second["rounds"] == 0
    assert second["holdings"] == 0
    assert second["values"] == 0
    assert second["policies"] == 0
    assert second["skipped_values"] == 4
    assert len(second["skipped_rounds"]) == 2
    assert len(store.list_rounds(db)) == 2
    assert len(store.list_holdings(db)) == 2
    with connect(db) as con:
        n = con.execute("SELECT COUNT(*) FROM holding_value").fetchone()[0]
    assert n == 4


def test_same_name_different_account_type_creates_two_holdings(tmp_path):
    """The real sheet has 'Schwab (Transfered from TD AmeriTrade)' twice:
    a Roth IRA and a separate Stock Account, both Steven's. Name alone must
    NOT be treated as the holding's identity, or the second row's values
    silently vanish (this is the $25,606.62 data-loss bug from cutover).
    """
    db = tmp_path / "t.db"
    init_db(db)
    xlsx = tmp_path / "dup_name_type.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Account", "Type", "Number", "Owner",
               "2026 Value (07-25-26)", "Notes"])
    ws.append(["Schwab (Transfered from TD AmeriTrade)", "Roth IRA Savings",
               "", "Steven", "$393.60", ""])
    ws.append(["Schwab (Transfered from TD AmeriTrade)", "Stock Account",
               "", "Steven", "$25,606.62", ""])
    wb.save(xlsx)

    result = imp.import_xlsx(db, xlsx)
    assert result["holdings"] == 2
    assert result["values"] == 2

    holdings = store.list_holdings(db)
    matching = [h for h in holdings
                if h["name"] == "Schwab (Transfered from TD AmeriTrade)"]
    assert len(matching) == 2
    assert {h["account_type"] for h in matching} == {
        "Roth IRA Savings", "Stock Account",
    }

    with connect(db) as con:
        cents = [
            r["value_cents"] for r in con.execute(
                "SELECT value_cents FROM holding_value WHERE holding_id IN (?, ?)",
                (matching[0]["id"], matching[1]["id"]),
            )
        ]
    assert sorted(cents) == [39360, 2560662]


def test_same_name_and_type_different_owner_creates_two_holdings(tmp_path):
    """'Treasury Direct - US Government' appears twice on the real sheet —
    same name, same account type, differing only by owner (Steven vs
    Allison). Owner must be part of the holding's identity too.
    """
    db = tmp_path / "t.db"
    init_db(db)
    xlsx = tmp_path / "dup_name_owner.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Account", "Type", "Number", "Owner",
               "2026 Value (07-25-26)", "Notes"])
    ws.append(["Treasury Direct - US Government", "Savings Bonds",
               "", "Steven", "$1,000.00", ""])
    ws.append(["Treasury Direct - US Government", "Savings Bonds",
               "", "Allison", "$2,000.00", ""])
    wb.save(xlsx)

    result = imp.import_xlsx(db, xlsx)
    assert result["holdings"] == 2
    assert result["values"] == 2

    holdings = store.list_holdings(db)
    matching = [h for h in holdings
                if h["name"] == "Treasury Direct - US Government"]
    assert len(matching) == 2
    assert {h["owner"] for h in matching} == {"Steven", "Allison"}

    with connect(db) as con:
        cents = [
            r["value_cents"] for r in con.execute(
                "SELECT value_cents FROM holding_value WHERE holding_id IN (?, ?)",
                (matching[0]["id"], matching[1]["id"]),
            )
        ]
    assert sorted(cents) == [100000, 200000]


def test_reimport_does_not_revert_a_hand_edited_value(tmp_path):
    db = tmp_path / "t.db"
    init_db(db)
    xlsx = _make_xlsx(tmp_path / "snap.xlsx")
    imp.import_xlsx(db, xlsx)

    rounds = {r["as_of_date"].isoformat(): r["id"] for r in store.list_rounds(db)}
    holdings = {h["name"]: h["id"] for h in store.list_holdings(db)}
    round_id = rounds["2026-07-25"]
    holding_id = holdings["Marcus"]

    # Operator hand-corrects the value in the UI — upsert_values stamps it
    # with source='manual'.
    store.upsert_values(
        db, round_id=round_id, source="manual",
        values=[{
            "holding_id": holding_id,
            "value_cents": 999900,
            "as_of_date": "2026-07-25",
        }],
    )

    imp.import_xlsx(db, xlsx)

    with connect(db) as con:
        row = con.execute(
            "SELECT value_cents, source FROM holding_value "
            "WHERE holding_id = ? AND round_id = ?",
            (holding_id, round_id),
        ).fetchone()
    assert row["value_cents"] == 999900
    assert row["source"] == "manual"


def test_roth_ira_is_classified_roth_not_pretax():
    assert imp._classify("Roth IRA") == ("retirement", "roth")
    assert imp._classify("Simple IRA") == ("retirement", "pretax")
    assert imp._classify("IRA") == ("retirement", "pretax")


def test_bare_year_header_raises_instead_of_inventing_jan_1(tmp_path):
    db = tmp_path / "t.db"
    init_db(db)
    xlsx = tmp_path / "bare_year.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Account", "Type", "Number", "Owner", "2026 Value", "Notes"])
    ws.append(["Marcus", "Savings", "1234", "Joint", "$1.00", ""])
    wb.save(xlsx)
    with pytest.raises(ValueError, match="no parseable date"):
        imp.import_xlsx(db, xlsx)


def test_property_rows_get_property_kind_and_detail(tmp_path):
    db = tmp_path / "t.db"
    init_db(db)
    imp.import_xlsx(db, _make_xlsx(tmp_path / "snap.xlsx"))
    holdings = {h["name"]: h for h in store.list_holdings(db)}
    assert holdings["117 Mayfield Dr"]["kind"] == "property"
    assert holdings["117 Mayfield Dr"]["is_primary_residence"] == 1
    assert holdings["Marcus"]["kind"] != "property"


def test_import_creates_policy_and_baseline_observation(tmp_path):
    db = tmp_path / "t.db"
    init_db(db)
    imp.import_xlsx(db, _make_xlsx(tmp_path / "snap.xlsx"))
    with connect(db) as con:
        policies = [dict(r) for r in con.execute("SELECT * FROM insurance_policy")]
        obs = [dict(r) for r in con.execute(
            "SELECT * FROM insurance_premium_observed"
        )]
    assert len(policies) == 1
    assert policies[0]["premium_cents"] == 279800
    assert len(obs) == 1
    assert obs[0]["amount_cents"] == 279800
    assert obs[0]["as_of_date"] == date(2026, 7, 25)   # newest round's date


def test_unparseable_header_raises(tmp_path):
    db = tmp_path / "t.db"
    init_db(db)
    xlsx = tmp_path / "bad.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Account", "Type", "Number", "Owner", "Value", "Notes"])
    ws.append(["Marcus", "Savings", "1234", "Joint", "$1.00", ""])
    wb.save(xlsx)
    with pytest.raises(ValueError, match="no parseable date"):
        imp.import_xlsx(db, xlsx)
